import os, sys, tempfile, unittest
from pathlib import Path

REPO = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(REPO / "scripts"))

import retirement_model as rm
import retirement_workbook as rw
from openpyxl import load_workbook


def sample_payload():
    assumptions = {
        "current_age": 37, "retirement_age": 62, "plan_through_age": 95,
        "current_balance": 500000.0, "annual_contribution": 40000.0,
        "retirement_spend": 120000.0, "inflation": 0.03,
        "expected_return": 0.075, "volatility": 0.13,
        "ss_start_age": 67, "ss_annual_amount": 40000.0,
        "scenarios": [{"name": "Retire at 60", "retirement_age": 60}],
    }
    result = rm.run_all(assumptions, n_paths=200, seed=1)
    return {
        "generated_date": "2026-07-23",
        "assumptions": assumptions,
        "assumption_sources": {
            "current_balance": "Monarch get_accounts (2026-07-23)",
            "annual_contribution": "Monarch get_cashflow trailing 12mo",
            "expected_return": "Derived from Monarch holdings (80/20)",
        },
        "result": result,
        "narrative": {
            "headline": "80% chance of funding retirement to 95.",
            "body": "At the current savings rate the plan succeeds in 4 of 5 simulated futures.",
        },
        "data_sources": [
            {"item": "Current balance", "value": "$500,000",
             "source": "Monarch get_accounts", "pulled": "2026-07-23", "overridden": "No"},
        ],
    }


class TestWorkbook(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False)
        self.tmp.close()
        self.path = self.tmp.name
        rw.build_workbook(sample_payload(), self.path)

    def tearDown(self):
        os.unlink(self.path)

    def test_expected_sheets_present(self):
        wb = load_workbook(self.path)
        for name in ["Summary", "Assumptions", "Base Case", "Monte Carlo", "Data Sources"]:
            self.assertIn(name, wb.sheetnames)
        # scenario sheet exists (name may be truncated/sanitized)
        self.assertTrue(any("Retire at 60" in s for s in wb.sheetnames))

    def test_base_case_uses_live_formulas_referencing_assumptions(self):
        wb = load_workbook(self.path)  # formulas kept as strings, not computed
        ws = wb["Base Case"]
        formula_cells = [
            c.value for row in ws.iter_rows() for c in row
            if isinstance(c.value, str) and c.value.startswith("=")
        ]
        self.assertTrue(formula_cells, "Base Case should contain live formulas")
        self.assertTrue(
            any("Assumptions" in f for f in formula_cells),
            "Base Case formulas should reference the Assumptions sheet",
        )

    def test_assumptions_sheet_has_editable_values(self):
        wb = load_workbook(self.path)
        ws = wb["Assumptions"]
        values = [c.value for row in ws.iter_rows() for c in row]
        self.assertIn(62, values)      # retirement age
        self.assertIn(500000.0, values)  # current balance


if __name__ == "__main__":
    unittest.main()
